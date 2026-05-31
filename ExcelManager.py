import pandas as pd
from openpyxl import load_workbook
import os
import shutil

class ExcelManager:
    def __init__(self):
        self.df = None
        self.file_path = None

    def load_excel(self, file_path):
        try:
            self.df = pd.read_excel(file_path, engine='openpyxl',dtype=str ).reset_index(drop=True)
            self.df = self.df.fillna(" ")
            self.file_path = file_path
            return True
        except Exception as e:
            print(f"Erro ao carregar arquivo Excel: {e}")
            return False
    
    def add_data(self, data):
        new_df = pd.DataFrame([data])
        if self.df is not None:
            if set(self.df.columns.tolist()) != set(new_df.columns.tolist()):
                print("Erro: As colunas do novo dado não correspondem às colunas existentes.")
                return False
            else:
                self.df = pd.concat([self.df, new_df], ignore_index=True)
                return True
        else:
            return False

    def save_excel(self):
        if self.df is not None:
            try:
                workbook = load_workbook(self.file_path)
                
                sheet = workbook.active
                if sheet.max_row > 1:
                    sheet.delete_rows(2, sheet.max_row - 1)
                    
                for row_idx, row in enumerate(self.df.values, start=2):
                    for col_idx, value in enumerate(row, start=1):
                        sheet.cell(row=row_idx, column=col_idx , value=value)
                        
                workbook.save(self.file_path)
                workbook.close()
                return True
            except Exception as e:
                print(f"Erro ao salvar arquivo Excel: {e}")
                return False
        else:
            
            return False

    def delete_data(self, index):
        if self.df is not None:
            try:
                self.df = self.df.drop(index)
                return True
            except Exception as e:
                print(f"Erro ao deletar dados: {e}")    
                return False
        else:
            return False

    def update_data(self, index, data):
        if self.df is not None:
            try:
                for key, value in data.items():
                    self.df.loc[index, key] = value
                return True
            except Exception as e:
                print(f"Erro ao atualizar dados: {e}")
                return False
        else:
            return False


    def get_columns(self):
        return self.df.columns.tolist()
    
    def save_backup(self):
        if not self.file_path:
            return False
        
        backup_dir = "backups"
        
        if not os.path.exists(backup_dir):
            os.makedirs(backup_dir)
            
        base_name = os.path.basename(self.file_path)
        name_without_ext, ext = os.path.splitext(base_name)
        backup_name = f"{name_without_ext}_backup{ext}"
        backup_path = os.path.join(backup_dir, backup_name)
        try:
            shutil.copy2(self.file_path, backup_path)
            return True
        except Exception as e:
            print(f"Erro ao criar backup: {e}")
            return False

